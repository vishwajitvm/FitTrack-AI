"""Exports: fitness_tracker.xlsx (6 sheets), workout_tracker.csv and auto file-sync (JSON backup: backup.py)."""
import csv
import io
import logging
from collections import defaultdict
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.db.session import SessionLocal
from app.models import DailySummary, FoodLog, User, WorkoutLog, WorkoutSession
from app.services.science import energy
from app.services.tracker import daily

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="0F172A")
HEADER_FONT = Font(color="FFFFFF", bold=True)

WORKOUT_HEADERS = [
    "ID", "Day", "Date", "Session ID", "Split", "Type", "Exercise", "Target Muscle", "Sets", "Reps", "Weight (kg)",
    "Cardio (min)", "Lifting (min)", "MET", "Calories", "RPE", "Notes",
]
FOOD_HEADERS = [
    "ID", "Day", "Date", "Time", "Meal", "Food", "Ingredients", "Quantity", "Unit", "Weight (g)", "Calories", "Protein (g)",
    "Carbs (g)", "Fat (g)", "Fibre (g)", "Calories Low", "Calories High", "Confidence", "Source", "Tags", "Notes",
]


def _num(v):
    """Blank for missing numbers, but keep a real 0 (e.g. a bodyweight exercise at 0 kg)."""
    return "" if v is None else v


async def _date_bounds(db: AsyncSession, user: User) -> tuple[date, date]:
    last = await db.scalar(select(func.max(DailySummary.date)))
    last_food = await db.scalar(select(func.max(FoodLog.date)))
    end = max(d for d in (last, last_food, user.start_date) if d is not None)
    return user.start_date, end


async def workout_rows(db: AsyncSession, user: User) -> list[list]:
    sessions = (await db.scalars(select(WorkoutSession).order_by(WorkoutSession.date, WorkoutSession.id))).all()
    logs = (await db.scalars(select(WorkoutLog).order_by(WorkoutLog.date, WorkoutLog.id))).all()
    split_by_session = {s.id: s.split for s in sessions}
    rows = []
    for s in sessions:
        rows.append([
            f"S{s.id}", daily.day_number(user, s.date), s.date.isoformat(), s.id, s.split, "lifting_block",
            f"{s.split.title()} session (lifting block)", "", "", "", "", "", s.lifting_duration_min, s.lifting_met,
            round(s.estimated_calories, 1), "", s.notes or "",
        ])
    for lg in logs:
        rows.append([
            lg.id, daily.day_number(user, lg.date), lg.date.isoformat(), lg.session_id or "",
            split_by_session.get(lg.session_id, "cardio" if lg.workout_type == "cardio" else ""), lg.workout_type,
            lg.exercise_name, lg.target_muscle or "", _num(lg.sets), lg.reps or "", _num(lg.weight_kg),
            _num(lg.cardio_duration_min), _num(lg.lifting_duration_min), _num(lg.met_value), round(lg.estimated_calories, 1),
            lg.rpe or "", lg.notes or "",
        ])
    rows.sort(key=lambda r: (r[2], 0 if r[5] == "cardio" else 1))
    return rows


def _write_sheet(ws, headers: list[str], rows: list[list]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for idx, header in enumerate(headers, start=1):
        width = max([len(str(header))] + [len(str(r[idx - 1])) for r in rows[:200] if idx - 1 < len(r)])
        ws.column_dimensions[get_column_letter(idx)].width = min(60, max(10, width + 2))


async def build_workbook(db: AsyncSession) -> bytes:
    user = await daily.get_user(db)
    start, end = await _date_bounds(db, user)
    days = await daily.compute_range(db, user, start, end)
    wb = Workbook()

    # 1. Daily Summary
    ws = wb.active
    ws.title = "Daily Summary"
    _write_sheet(ws, [
        "Day", "Date", "Weekday", "Morning Weight (kg)", "Calories", "Protein (g)", "Carbs (g)", "Fat (g)", "Fibre (g)",
        "Whey Scoops", "Protein g/kg", "Exercise kcal", "BMR", "NEAT", "TEF", "TDEE", "Deficit", "Readiness", "Energy",
        "Sleep (h)", "Bloating", "Alerts", "Coach Advice",
    ], [[
        d["day_number"], d["date"], d["weekday"], d["morning_weight"], d["total_calories"], d["total_protein"], d["total_carbs"],
        d["total_fat"], d["total_fibre"], d["whey_scoops"], d["protein_per_kg"], d["estimated_exercise_calories"],
        d["estimated_bmr"] if d["estimated_bmr"] is not None else "Pending", d["estimated_neat"], d["estimated_tef"],
        d["estimated_tdee"] if d["estimated_tdee"] is not None else "Pending",
        d["calorie_deficit"] if d["calorie_deficit"] is not None else "Pending Biometric Confirmation",
        d["readiness_score"], d["energy_level"], d["sleep_hours"], d["bloating_status"],
        "; ".join(a["title"] for a in d["alerts"] if a["level"] in ("critical", "warning")), d["coach_advice"],
    ] for d in days])

    # 2. Food Log
    foods = (await db.scalars(select(FoodLog).order_by(FoodLog.date, FoodLog.time, FoodLog.id))).all()
    _write_sheet(wb.create_sheet("Food Log"), FOOD_HEADERS, [[
        f.id, daily.day_number(user, f.date), f.date.isoformat(), f.time or "", f.meal_type, f.food_name,
        ", ".join(f.ingredients_json or []), f.quantity, f.portion_unit or "", f.weight_grams, f.calories, f.protein, f.carbs,
        f.fat, f.fibre, f.calories_low, f.calories_high, f.confidence_score, f.research_source or "", ", ".join(f.tags or []),
        f.notes or "",
    ] for f in foods])

    # 3. Workout Log
    _write_sheet(wb.create_sheet("Workout Log"), WORKOUT_HEADERS, await workout_rows(db, user))

    # 4. Body Weight
    weights = [(d["day_number"], d["date"], d["morning_weight"]) for d in days if d["morning_weight"]]
    rows = []
    for i, (num, dt, w) in enumerate(weights):
        window = [x[2] for x in weights[max(0, i - 6) : i + 1]]
        rows.append([num, dt, w, round(w - user.starting_weight, 2), round(sum(window) / len(window), 2)])
    _write_sheet(wb.create_sheet("Body Weight"), ["Day", "Date", "Morning Weight (kg)", "Change from Start (kg)", "7-day Avg (kg)"], rows)

    # 5. Nutrition Targets
    weight = user.current_weight or user.starting_weight
    bmr = energy.bmr_mifflin(weight, user.height_cm, user.age, user.sex)
    pt = energy.protein_target(weight, user.protein_target_g_per_kg)
    _write_sheet(wb.create_sheet("Nutrition Targets"), ["Metric", "Value", "Notes"], [
        ["Current weight (kg)", weight, "Latest morning weight"],
        ["Height (cm)", user.height_cm, ""],
        ["Age / Sex", f"{user.age or 'pending'} / {user.sex or 'pending'}", "Needed for BMR"],
        ["BMR (Mifflin-St Jeor)", round(bmr) if bmr else "Pending Biometric Confirmation", "10W + 6.25H − 5A ± s"],
        ["Protein target (g/day)", f"{pt['min']}–{pt['max']} (target {pt['target']})", "1.6–2.0 g/kg"],
        ["Fibre target (g/day)", "≥ 25–35", "Gut health & satiety"],
        ["Water target (ml/day)", user.water_target_ml or round(weight * 35), "~35 ml/kg"],
        ["Recommended deficit", "500–750 kcal/day", "≈ 0.5–0.75 kg/week"],
        ["Crash threshold", "< 1,000 kcal intake", "Critical alert"],
        ["Rest day", (user.rest_day or "sunday").title(), "Mandatory rest & gut recovery"],
        ["Gym return / Bedtime", f"{user.gym_return_time or '–'} / {user.bedtime or '–'}", "Main meal before gym; light protein after"],
    ])

    # 6. Weekly Report
    weeks: dict[date, list[dict]] = defaultdict(list)
    for d in days:
        dd = date.fromisoformat(d["date"])
        weeks[dd - timedelta(days=dd.weekday())].append(d)
    weekly = []
    for wk, items in sorted(weeks.items()):
        logged = [d for d in items if d["food_count"]]
        n = len(logged) or 1
        deficits = [d["calorie_deficit"] for d in logged if d["calorie_deficit"] is not None]
        ws_weights = [d["morning_weight"] for d in items if d["morning_weight"]]
        weekly.append([
            wk.isoformat(), len(logged), round(sum(d["total_calories"] for d in logged) / n),
            round(sum(d["total_protein"] for d in logged) / n, 1), round(sum(deficits) / len(deficits)) if deficits else "Pending",
            sum(d["estimated_exercise_calories"] for d in items), sum(1 for d in items if d["has_workout"]),
            ws_weights[0] if ws_weights else "", ws_weights[-1] if ws_weights else "",
            sum(1 for d in items for a in d["alerts"] if a["level"] == "critical"),
        ])
    _write_sheet(wb.create_sheet("Weekly Report"), [
        "Week Starting", "Days Logged", "Avg Calories", "Avg Protein (g)", "Avg Deficit", "Exercise kcal", "Workouts",
        "Weight Start", "Weight End", "Critical Alerts",
    ], weekly)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def build_workout_csv(db: AsyncSession) -> str:
    user = await daily.get_user(db)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(WORKOUT_HEADERS)
    writer.writerows(await workout_rows(db, user))
    return buf.getvalue()


async def sync_files() -> list[str]:
    """Writes fitness_tracker.xlsx and workout_tracker.csv into EXPORT_DIR."""
    out_dir = settings.export_dir
    out_dir.mkdir(parents=True, exist_ok=True)
    try:
        async with SessionLocal() as db:
            (out_dir / "fitness_tracker.xlsx").write_bytes(await build_workbook(db))
            (out_dir / "workout_tracker.csv").write_text(await build_workout_csv(db), encoding="utf-8")
    except Exception as exc:  # noqa: BLE001  file sync must never break an API call
        log.warning("File sync failed: %s", exc)
        return []
    return [str(out_dir / "fitness_tracker.xlsx"), str(out_dir / "workout_tracker.csv")]


async def sync_if_enabled() -> None:
    if settings.auto_sync_files:
        await sync_files()
