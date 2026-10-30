"""Unit tests for Excel import cell parsing and JSON backup validation (no database needed).

Run:  pytest tests/test_data_portability.py
"""
import io
import json
from datetime import date, datetime, time

import pytest
from openpyxl import Workbook

from app.services.data import backup, importer
from app.services.data.importer import MEAL_ALIASES, MEALS, SPLIT_ALIASES, SPLITS, RowError


def _xlsx(sheets: dict[str, list[list]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_ids_from_excel():
    assert importer._id(12) == 12
    assert importer._id(12.0) == 12
    assert importer._id("12") == 12
    assert importer._id(" 12.0 ") == 12
    assert importer._id("S7") == 7
    for bad in (None, "", "abc", 0, -3, 1.5, "S"):
        assert importer._id(bad) is None


def test_dates_from_excel():
    assert importer._d("2026-10-26") == date(2026, 10, 26)
    assert importer._d("2026-10-26 00:00:00") == date(2026, 10, 26)
    assert importer._d(datetime(2026, 10, 26, 7, 30)) == date(2026, 10, 26)
    assert importer._d(date(2026, 10, 26)) == date(2026, 10, 26)
    assert importer._d(46321) == date(2026, 10, 26)  # Excel serial day number
    assert importer._d("26/10/2026") == date(2026, 10, 26)
    assert importer._d("26 Oct 2026") == date(2026, 10, 26)
    assert importer._d(None) is None and importer._d("  ") is None
    for bad in ("yesterday", "2026-13-45", 5, True):
        with pytest.raises(RowError):
            importer._d(bad)


def test_times_from_excel():
    assert importer._t("06:36") == "06:36"
    assert importer._t("6:36") == "06:36"
    assert importer._t("06:36:00") == "06:36"
    assert importer._t("6.36") == "06:36"
    assert importer._t("6:36 PM") == "18:36"
    assert importer._t("12:05 am") == "00:05"
    assert importer._t(time(21, 10)) == "21:10"
    assert importer._t(datetime(2026, 10, 26, 9, 5)) == "09:05"
    assert importer._t(0.5) == "12:00"  # Excel stores times as a fraction of a day
    assert importer._t(None) is None
    for bad in ("25:00", "noon", "13:00 PM", 2.5):
        with pytest.raises(RowError):
            importer._t(bad)


def test_numbers_and_text():
    assert importer._f("1,279") == 1279.0
    assert importer._f(0) == 0.0  # a real zero is kept
    assert importer._f("") is None
    with pytest.raises(RowError):
        importer._f("lots")
    with pytest.raises(RowError):
        importer._nonneg(-5, "Calories")
    assert importer._s(12.0) == "12"
    assert importer._s("  Paneer  ") == "Paneer"
    assert importer._s("x" * 300, 200) == "x" * 200
    assert importer._s(datetime(2026, 10, 15)) == "2026-10-15"


def test_meal_and_split_names_typed_in_excel():
    choice = importer._choice
    assert choice("Lunch", MEALS, MEAL_ALIASES, "other") == "lunch"
    assert choice("Morning Snack", MEALS, MEAL_ALIASES, "other") == "morning_snack"
    assert choice("post-workout", MEALS, MEAL_ALIASES, "other") == "post_workout"
    assert choice("Snack", MEALS, MEAL_ALIASES, "other") == "evening_snack"
    assert choice("brunch", MEALS, MEAL_ALIASES, "other") == "other"
    assert choice(None, MEALS, MEAL_ALIASES, "other") == "other"
    assert choice("Pull", SPLITS, SPLIT_ALIASES, "other") == "pull"
    assert choice("Full Body", SPLITS, SPLIT_ALIASES, "other") == "full_body"
    assert choice("Leg", SPLITS, SPLIT_ALIASES, "other") == "legs"


def test_file_checks():
    with pytest.raises(importer.ImportError_, match="empty"):
        importer._open(b"")
    with pytest.raises(importer.ImportError_, match="isn't an Excel"):
        importer._open(b'{"format": "fittrack-backup"}')
    with pytest.raises(importer.ImportError_, match="FitTrack export"):
        importer._open(_xlsx({"Sheet1": [["a", "b"], [1, 2]]}))
    wb = importer._open(_xlsx({"Food Log": [["ID", "Date", "Food"], [1, "2026-10-26", "Roti"]]}))
    assert "Food Log" in wb.sheetnames


def test_same_entry_guard():
    assert importer._same(date(2026, 10, 26), "Roti", date(2026, 10, 26), "Something else")  # same day
    assert importer._same(date(2026, 10, 20), "Roti", date(2026, 10, 26), "roti")  # date edited, same food
    assert not importer._same(date(2026, 10, 20), "Roti", date(2026, 10, 26), "Paneer")  # an unrelated row


def test_backup_parse_rejects_bad_files():
    for data, msg in [
        (b"", "empty"),
        (b"not json", "valid JSON"),
        (b"[1, 2]", "isn't a FitTrack backup"),
        (b'{"hello": 1}', "isn't a FitTrack backup"),
        (json.dumps({"format": "fittrack-backup", "version": 99, "tables": {}}).encode(), "newer"),
        (json.dumps({"format": "fittrack-backup", "version": 1, "tables": {"users": []}}).encode(), "no profile"),
        (json.dumps({"format": "fittrack-backup", "version": 1, "tables": {"users": [{"id": 1}], "food_logs": "x"}}).encode(), "damaged"),
    ]:
        with pytest.raises(backup.BackupError, match=msg):
            backup.parse(data)


def test_backup_parse_accepts_new_and_old_formats():
    new = backup.parse(json.dumps({"format": "fittrack-backup", "version": 1, "tables": {"users": [{"id": 1}]}}).encode("utf-8-sig"))
    assert new["format"] == "fittrack-backup"
    assert backup.preview(new)["mode"] == "replace"
    old = backup.parse(json.dumps({"profile": {}, "food_logs": [{"id": 1}], "workout_logs": []}).encode())
    assert old["format"] == "legacy"
    assert backup.preview(old) == {"kind": "legacy", "mode": "merge", "created_at": None, "counts": {"food_logs": 1, "workout_sessions": 0, "workout_logs": 0, "bio_feedback": 0}, "total": 1}


def test_backup_values_are_type_checked():
    from app.models import FoodLog

    cols = FoodLog.__table__.c
    assert backup._from_json(cols.date, "2026-10-26") == date(2026, 10, 26)
    assert backup._from_json(cols.calories, 12) == 12.0
    assert backup._from_json(cols.id, 5.0) == 5
    for col, value in [(cols.calories, "lots"), (cols.id, 1.5), (cols.date, "soon"), (cols.time, "123456"), (cols.calories, True)]:
        with pytest.raises(backup.BackupError):
            backup._from_json(col, value)
