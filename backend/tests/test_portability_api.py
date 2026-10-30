"""End-to-end scenarios for Excel import and JSON backup/restore against a RUNNING backend.

These tests write data, so they only run when FITTRACK_TEST_API points at a throwaway test stack — never at the
backend holding your real data. Each test restores the starting data from a JSON backup when it finishes.

Run:  FITTRACK_TEST_API=http://localhost:8001 pytest tests/test_portability_api.py
"""
import io
import json
import os
import zipfile
from datetime import datetime, time

import pytest

httpx = pytest.importorskip("httpx")
from openpyxl import Workbook, load_workbook  # noqa: E402

API = os.environ.get("FITTRACK_TEST_API", "").rstrip("/")
pytestmark = pytest.mark.skipif(not API, reason="set FITTRACK_TEST_API to a throwaway test backend to run these")
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
VOLATILE = {"updated_at", "created_at"}


@pytest.fixture(scope="module")
def client():
    with httpx.Client(base_url=API, timeout=60) as c:
        assert c.get("/api/health").json()["status"] == "ok"
        yield c


@pytest.fixture(autouse=True)
def restore_after(client):
    original = client.get("/api/export/json").content
    yield original
    r = client.post("/api/import/json", files={"file": ("start.json", original, "application/json")})
    assert r.status_code == 200, r.text


def backup(client) -> dict:
    return json.loads(client.get("/api/export/json").content)


def rows(data: dict, table: str) -> dict[int, dict]:
    return {r["id"]: {k: v for k, v in r.items() if k not in VOLATILE} for r in data["tables"][table]}


def upload_xlsx(client, data: bytes, name="fitness_tracker.xlsx"):
    return client.post("/api/import/xlsx", files={"file": (name, data, XLSX)})


def save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def sheet_rows(ws) -> tuple[list[str], list[list]]:
    values = [list(r) for r in ws.iter_rows(values_only=True)]
    return [str(h) for h in values[0]], values[1:]


def smoke(client):
    for path in ("/api/health", "/api/dashboard", "/api/today", "/api/food", "/api/workouts", "/api/history", "/api/profile"):
        r = client.get(path)
        assert r.status_code in (200, 404, 405), f"{path} → {r.status_code}"
        assert r.status_code != 500, f"{path} crashed"


def test_unchanged_round_trip_changes_nothing(client):
    before = backup(client)
    r = upload_xlsx(client, client.get("/api/export/xlsx").content)
    assert r.status_code == 200, r.text
    assert r.json()["skipped"] == 0 and r.json()["created"] == {"food": 0, "workouts": 0}
    after = backup(client)
    for table in ("food_logs", "workout_logs", "workout_sessions"):
        assert rows(before, table) == rows(after, table), f"{table} changed on an unchanged re-import"
    smoke(client)


def test_round_trip_twice_does_not_duplicate(client):
    first = backup(client)["counts"]
    for _ in range(2):
        assert upload_xlsx(client, client.get("/api/export/xlsx").content).status_code == 200
    assert backup(client)["counts"] == first


def test_excel_style_edits(client):
    wb = load_workbook(io.BytesIO(client.get("/api/export/xlsx").content))
    ws = wb["Food Log"]
    headers, data = sheet_rows(ws)
    col = {h: i + 1 for i, h in enumerate(headers)}
    target = data[0]
    food_id = int(target[0])
    # what Excel does to a file: IDs as floats, real dates/times, typed words in any case
    ws.cell(2, col["ID"], float(food_id))
    ws.cell(2, col["Date"], datetime.fromisoformat(str(target[col["Date"] - 1])))
    ws.cell(2, col["Time"], time(7, 45))
    ws.cell(2, col["Meal"], "Lunch")
    ws.cell(2, col["Calories"], "321")
    new = [None] * len(headers)
    new[col["Date"] - 1], new[col["Food"] - 1], new[col["Meal"] - 1], new[col["Calories"] - 1], new[col["Time"] - 1] = "2026-10-27", "Test banana", "Morning Snack", 105, "6:30 PM"
    ws.append(new)
    bad = list(new)
    bad[col["Food"] - 1], bad[col["Calories"] - 1] = "Broken row", -50
    ws.append(bad)
    bad_date = list(new)
    bad_date[col["Food"] - 1], bad_date[col["Date"] - 1] = "No date row", "someday"
    ws.append(bad_date)
    r = upload_xlsx(client, save(wb))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["created"]["food"] == 1
    assert body["skipped"] == 2
    assert {e["row"] for e in body["errors"]} == {ws.max_row - 1, ws.max_row}
    foods = rows(backup(client), "food_logs")
    edited = foods[food_id]
    assert (edited["time"], edited["meal_type"], edited["calories"]) == ("07:45", "lunch", 321.0)
    banana = next(f for f in foods.values() if f["food_name"] == "Test banana")
    assert (banana["meal_type"], banana["time"], banana["date"]) == ("morning_snack", "18:30", "2026-10-27")
    assert not any(f["food_name"] in ("Broken row", "No date row") for f in foods.values())
    smoke(client)


def test_excel_workout_edits_keep_session_links(client):
    wb = load_workbook(io.BytesIO(client.get("/api/export/xlsx").content))
    ws = wb["Workout Log"]
    headers, data = sheet_rows(ws)
    col = {h: i + 1 for i, h in enumerate(headers)}
    for i, r in enumerate(data, start=2):
        if r[col["Session ID"] - 1] not in (None, ""):
            ws.cell(i, col["Session ID"], float(r[col["Session ID"] - 1]))  # Excel float
        if r[col["Type"] - 1] == "lifting_block":
            ws.cell(i, col["Split"], str(r[col["Split"] - 1]).title())  # "Pull"
    before = backup(client)
    assert upload_xlsx(client, save(wb)).status_code == 200
    after = backup(client)
    assert rows(before, "workout_logs") == rows(after, "workout_logs")
    assert rows(before, "workout_sessions") == rows(after, "workout_sessions")


def test_workbook_from_another_install_does_not_overwrite(client):
    data = backup(client)
    some = data["tables"]["food_logs"][0]
    wb = Workbook()
    ws = wb.active
    ws.title = "Food Log"
    ws.append(["ID", "Date", "Meal", "Food", "Calories"])
    ws.append([some["id"], "2026-01-05", "dinner", "Someone else's pizza", 900])  # same id, different entry
    r = upload_xlsx(client, save(wb))
    assert r.status_code == 200, r.text
    assert r.json()["created"]["food"] == 1
    foods = rows(backup(client), "food_logs")
    assert foods[some["id"]]["food_name"] == some["food_name"]
    assert any(f["food_name"] == "Someone else's pizza" for f in foods.values())


@pytest.mark.parametrize(
    ("name", "content", "status", "message"),
    [
        ("backup.json", b"{}", 400, "JSON backup"),
        ("notes.txt", b"hello", 400, ".xlsx or .xlsm"),
        ("empty.xlsx", b"", 400, "empty"),
        ("fake.xlsx", b"PK\x03\x04 not really a zip", 400, "isn't an Excel"),
        ("other.xlsx", None, 400, "FitTrack export"),
    ],
)
def test_bad_excel_files_are_rejected_cleanly(client, name, content, status, message):
    if content is None:
        wb = Workbook()
        wb.active.append(["Name", "Value"])
        content = save(wb)
    before = backup(client)["counts"]
    r = upload_xlsx(client, content, name)
    assert r.status_code == status, r.text
    assert message in r.json()["detail"]
    assert backup(client)["counts"] == before
    smoke(client)


def test_damaged_zip_with_valid_name(client):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "not a workbook")
    r = upload_xlsx(client, buf.getvalue())
    assert r.status_code == 400 and "couldn't be read" in r.json()["detail"]


def test_xlsm_is_accepted(client):
    r = upload_xlsx(client, client.get("/api/export/xlsx").content, "fitness_tracker.xlsm")
    assert r.status_code == 200, r.text


def test_json_restore_round_trip_is_exact(client, restore_after):
    before = backup(client)
    r = client.post("/api/import/json", files={"file": ("b.json", restore_after, "application/json")})
    assert r.status_code == 200, r.text
    assert r.json()["mode"] == "replace" and r.json()["safety_backup"]
    after = backup(client)
    for table in before["tables"]:
        assert before["tables"][table] == after["tables"][table], table


def test_json_restore_replaces_and_can_be_undone(client, restore_after):
    added = client.post("/api/food", json={"date": "2026-10-27", "meal_type": "dinner", "food_name": "Restore test dal", "calories": 200, "protein": 10})
    assert added.status_code == 200, added.text
    with_dal = client.get("/api/export/json").content
    # restore the starting backup: the dal disappears
    r = client.post("/api/import/json", files={"file": ("b.json", restore_after, "application/json")})
    assert r.status_code == 200
    assert not any(f["food_name"] == "Restore test dal" for f in backup(client)["tables"]["food_logs"])
    # the safety copy taken before the restore has it, and can be downloaded and restored
    safety = r.json()["safety_backup"]
    assert safety in [b["name"] for b in client.get("/api/backups").json()]
    copy = client.get(f"/api/backups/{safety}")
    assert copy.status_code == 200
    assert json.loads(copy.content)["tables"]["food_logs"] == json.loads(with_dal)["tables"]["food_logs"]
    r = client.post("/api/import/json", files={"file": (safety, copy.content, "application/json")})
    assert r.status_code == 200
    assert any(f["food_name"] == "Restore test dal" for f in backup(client)["tables"]["food_logs"])
    # new entries after a restore get fresh ids (sequences were reset)
    again = client.post("/api/food", json={"date": "2026-10-27", "meal_type": "dinner", "food_name": "After restore", "calories": 50, "protein": 1})
    assert again.status_code == 200, again.text
    smoke(client)


def test_bad_json_backups_change_nothing(client, restore_after):
    before = backup(client)
    good = json.loads(restore_after)
    broken_value = json.loads(restore_after)
    broken_value["tables"]["food_logs"][0]["calories"] = "lots"
    too_long = json.loads(restore_after)
    too_long["tables"]["food_logs"][0]["meal_type"] = "x" * 50
    dangling = json.loads(restore_after)
    dangling["tables"]["food_logs"][0]["daily_summary_id"] = 999999  # foreign key that doesn't exist
    cases = [
        (b"", "empty"),
        (b"{not json", "valid JSON"),
        (json.dumps({"format": "fittrack-backup", "version": 1, "tables": {}}).encode(), "no profile"),
        (json.dumps({**good, "version": 7}).encode(), "newer"),
        (json.dumps(broken_value).encode(), "calories"),
        (json.dumps(too_long).encode(), "meal_type"),
        (json.dumps(dangling).encode(), "nothing was changed"),
    ]
    for content, message in cases:
        r = client.post("/api/import/json", files={"file": ("b.json", content, "application/json")})
        assert r.status_code == 400, (message, r.text)
        assert message in r.json()["detail"], (message, r.json()["detail"])
        assert backup(client)["tables"] == before["tables"], f"data changed after a rejected file ({message})"
    r = client.post("/api/import/json", files={"file": ("x.xlsx", b"PK", XLSX)})
    assert r.status_code == 400 and "Excel" in r.json()["detail"]
    smoke(client)


def test_json_preview(client, restore_after):
    r = client.post("/api/import/json/preview", files={"file": ("b.json", restore_after, "application/json")})
    assert r.status_code == 200
    body = r.json()
    assert body["mode"] == "replace" and body["counts"]["food_logs"] == len(json.loads(restore_after)["tables"]["food_logs"])


def test_old_json_format_merges(client, restore_after):
    data = json.loads(restore_after)
    food = data["tables"]["food_logs"][0]
    legacy = {
        "profile": {},
        "food_logs": [
            {**{k: food[k] for k in ("id", "date", "meal_type", "time", "food_name", "calories", "protein")}, "calories": 999, "ingredients": []},
            {"id": None, "date": "2026-10-28", "meal_type": "dinner", "food_name": "Legacy khichdi", "calories": 300, "protein": 9, "ingredients": ["rice", "dal"]},
        ],
        "workout_sessions": [],
        "workout_logs": [],
        "bio_feedback": [],
    }
    r = client.post("/api/import/json", files={"file": ("old.json", json.dumps(legacy).encode(), "application/json")})
    assert r.status_code == 200, r.text
    assert r.json()["mode"] == "merge"
    foods = rows(backup(client), "food_logs")
    assert foods[food["id"]]["calories"] == 999
    khichdi = next(f for f in foods.values() if f["food_name"] == "Legacy khichdi")
    assert khichdi["ingredients_json"] == ["rice", "dal"]
    assert len(foods) == len(data["tables"]["food_logs"]) + 1
